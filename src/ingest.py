"""Parse EDRM-specific files into JSONL documents and chunks.

Default source is data/extracted/EDRM specific data (not WebKB HTML).

Run from the repo root:
    C:\\Users\\gokif\\anaconda3\\python.exe src\\ingest.py
"""

from __future__ import annotations

import argparse
import hashlib
import json
import mailbox
import re
import zipfile
from email.message import Message
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
EXTRACTED = ROOT / "data" / "extracted"
PROCESSED = ROOT / "data" / "processed"
DOCS_OUT = PROCESSED / "documents.jsonl"
CHUNKS_OUT = PROCESSED / "chunks.jsonl"

INCLUDE = {
    ".txt",
    ".html",
    ".htm",
    ".xml",
    ".xlsx",
    ".xltx",
    ".pptx",
    ".docx",
    ".docm",
    ".dotm",
    ".pdf",
    ".csv",
    ".mbox",
}
SKIP = {
    ".jpg",
    ".jpeg",
    ".tif",
    ".tiff",
    ".mp3",
    ".doc",
    ".rtf",
    ".odt",
    ".pot",
    ".pps",
    ".xla",
    ".dwg",
    ".dxf",
    ".dmg",
    ".snoop",
    ".zip",
    ".gz",
}

CHUNK_CHARS = 1200
CHUNK_OVERLAP = 200
SS = "{urn:schemas-microsoft-com:office:spreadsheet}"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def chunk_text(text: str) -> list[str]:
    text = clean_text(text)
    if not text:
        return []
    if len(text) <= CHUNK_CHARS:
        return [text]
    chunks = []
    start = 0
    while start < len(text):
        end = min(len(text), start + CHUNK_CHARS)
        if end < len(text):
            split = text.rfind("\n", start + CHUNK_CHARS // 2, end)
            if split == -1:
                split = text.rfind(" ", start + CHUNK_CHARS // 2, end)
            if split != -1:
                end = split
        piece = text[start:end].strip()
        if piece:
            chunks.append(piece)
        if end >= len(text):
            break
        start = max(end - CHUNK_OVERLAP, start + 1)
    return chunks


def html_to_text(raw: str) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(raw, "lxml")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    return soup.get_text("\n", strip = True)


def parse_txt(path: Path) -> str:
    data = path.read_bytes()
    for enc in ("utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors = "replace")


def parse_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only = True, read_only = True)
    parts = []
    for sheet in wb.worksheets:
        parts.append(f"# sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only = True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cell.strip() for cell in cells):
                parts.append("\t".join(cells))
    wb.close()
    return "\n".join(parts)


def parse_spreadsheet_xml(path: Path) -> str:
    tree = ET.parse(path)
    parts = []
    for worksheet in tree.findall(f".//{SS}Worksheet"):
        name = worksheet.attrib.get(f"{SS}Name", "sheet")
        parts.append(f"# sheet: {name}")
        for row in worksheet.findall(f".//{SS}Row"):
            cells = []
            for cell in row.findall(f"{SS}Cell"):
                data = cell.find(f"{SS}Data")
                cells.append("" if data is None or data.text is None else data.text)
            if any(c.strip() for c in cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


def parse_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, start = 1):
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                texts.append(shape.text)
        body = clean_text("\n".join(texts))
        if body:
            parts.append(f"# slide {i}\n{body}")
    return "\n\n".join(parts)


def parse_docx(path: Path) -> str:
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def parse_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def message_body(msg: Message) -> str:
    if msg.is_multipart():
        parts = []
        for part in msg.walk():
            ctype = part.get_content_type()
            if ctype == "text/plain":
                payload = part.get_payload(decode=True) or b""
                charset = part.get_content_charset() or "utf-8"
                parts.append(payload.decode(charset, errors = "replace"))
            elif ctype == "text/html" and not parts:
                payload = part.get_payload(decode = True) or b""
                charset = part.get_content_charset() or "utf-8"
                parts.append(html_to_text(payload.decode(charset, errors = "replace")))
        return "\n".join(parts)
    payload = msg.get_payload(decode = True)
    if payload is None:
        return str(msg.get_payload() or "")
    charset = msg.get_content_charset() or "utf-8"
    text = payload.decode(charset, errors = "replace")
    if msg.get_content_type() == "text/html":
        return html_to_text(text)
    return text


def attachment_names(msg: Message) -> list[str]:
    names = []
    for part in msg.walk():
        name = part.get_filename()
        if name:
            names.append(name)
    return names


def parse_file(path: Path) -> list[dict]:
    ext = path.suffix.lower()
    rel = str(path.relative_to(EXTRACTED)) if path.is_relative_to(EXTRACTED) else str(path)
    file_hash = sha256_file(path)
    base = {
        "source_path": rel,
        "file_type": ext,
        "sha256": file_hash,
        "bytes": path.stat().st_size,
        "title": path.name,
    }

    if ext == ".mbox":
        docs = []
        for i, msg in enumerate(mailbox.mbox(path)):
            body = clean_text(message_body(msg))
            subject = msg.get("subject") or "(no subject)"
            docs.append(
                {
                    **base,
                    "doc_id": f"{file_hash}:msg:{i}",
                    "title": subject,
                    "text": body,
                    "metadata": {
                        "from": msg.get("from"),
                        "to": msg.get("to"),
                        "date": msg.get("date"),
                        "subject": subject,
                        "message_id": msg.get("message-id"),
                        "attachments": attachment_names(msg),
                    },
                }
            )
        return docs

    if ext in {".txt", ".csv"}:
        text = parse_txt(path)
    elif ext in {".html", ".htm"}:
        text = html_to_text(parse_txt(path))
    elif ext == ".xml":
        text = parse_spreadsheet_xml(path)
    elif ext in {".xlsx", ".xltx"}:
        text = parse_xlsx(path)
    elif ext == ".pptx":
        text = parse_pptx(path)
    elif ext in {".docx", ".docm", ".dotm"}:
        text = parse_docx(path)
    elif ext == ".pdf":
        text = parse_pdf(path)
    else:
        return []

    return [
        {
            **base,
            "doc_id": file_hash,
            "text": clean_text(text),
            "metadata": {},
        }
    ]


def is_ooxml(path: Path) -> bool:
    try:
        return zipfile.is_zipfile(path)
    except OSError:
        return False


def should_parse(path: Path) -> bool:
    ext = path.suffix.lower()
    if ext in SKIP or ext not in INCLUDE:
        return False
    if ext in {".docx", ".docm", ".dotm", ".pptx", ".xlsx", ".xltx"}:
        return is_ooxml(path)
    return True


def iter_source_files(source: Path) -> list[Path]:
    return sorted(p for p in source.rglob("*") if p.is_file() and should_parse(p))


def main() -> None:
    parser = argparse.ArgumentParser(
        description = "Parse extracted documents into JSONL documents and chunks."
    )
    parser.add_argument(
        "--source",
        default = "EDRM specific data",
        help = "Source directory relative to data/extracted",
    )
    parser.add_argument(
        "--corpus",
        default = "edrm",
        help = "Name used for the processed output directory",
    )
    args = parser.parse_args()

    source = EXTRACTED / args.source

    if not source.exists():
        raise SystemExit(f"missing source: {source}\nrun extract.py first")

    corpus_dir = PROCESSED / args.corpus
    corpus_dir.mkdir(parents = True, exist_ok = True)

    docs_out = corpus_dir / "documents.jsonl"
    chunks_out = corpus_dir / "chunks.jsonl"
    files = iter_source_files(source)
    docs = []
    errors = []
    for path in files:
        try:
            parsed = parse_file(path)
        except Exception as exc:  # noqa: BLE001 — keep ingest going per file
            errors.append({"path": str(path), "error": str(exc)})
            print(f"skip {path.name}: {exc}")
            continue
        for doc in parsed:
            if doc.get("text"):
                docs.append(doc)
            else:
                errors.append({"path": str(path), "error": "empty text"})

    chunks = []
    for doc in docs:
        pieces = chunk_text(doc["text"])
        for i, piece in enumerate(pieces):
            chunks.append(
                {
                    "chunk_id": f"{doc['doc_id']}:{i}",
                    "doc_id": doc["doc_id"],
                    "source_path": doc["source_path"],
                    "file_type": doc["file_type"],
                    "title": doc["title"],
                    "chunk_index": i,
                    "text": piece,
                    "metadata": doc.get("metadata") or {},
                }
            )

    with docs_out.open("w", encoding = "utf-8") as fh:
        for doc in docs:
            row = {k: v for k, v in doc.items() if k != "text"}
            row["chars"] = len(doc["text"])
            fh.write(json.dumps(row, ensure_ascii=False) + "\n")

    with chunks_out.open("w", encoding = "utf-8") as fh:
        for chunk in chunks:
            fh.write(json.dumps(chunk, ensure_ascii = False) + "\n")

    print(f"source: {source}")
    print(f"parsed files: {len(files)}")
    print(f"documents: {len(docs)} -> {docs_out}")
    print(f"chunks: {len(chunks)} -> {chunks_out}")
    print(f"errors: {len(errors)}")
    by_type: dict[str, int] = {}
    for doc in docs:
        by_type[doc["file_type"]] = by_type.get(doc["file_type"], 0) + 1
    print("documents by type:", json.dumps(by_type, indent = 2))


if __name__ == "__main__":
    main()
